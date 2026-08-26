from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "artifacts"
OUTPUT_DIR.mkdir(exist_ok=True)
OUTPUT = OUTPUT_DIR / "Reglamento_de_Uso_del_Sistema_Certificador_INSO.docx"
LOGO = ROOT / "public" / "images" / "branding" / "inso-logo.png"

NAVY = "123B2A"
GREEN = "08783E"
MID_GREEN = "4F9F6E"
LIGHT_GREEN = "EAF5EE"
PALE = "F5F8F6"
GRAY = "5F6B66"
LIGHT_GRAY = "E6ECE8"
BLACK = "1B1F1D"
WHITE = "FFFFFF"
GOLD = "D99A22"


def set_run_font(run, name="Calibri", size=None, color=BLACK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + side
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B9C8BF", size="5"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
    set_table_borders(table)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, sep, text, end])
    return run


def add_hyperlink(paragraph, text, url, color=GREEN):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    new_run.append(r_pr)
    node = OxmlElement("w:t")
    node.text = text
    new_run.append(node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_num_definition(doc, num_id, abstract_id, fmt, text, left=720, hanging=360, font="Calibri"):
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    lvl.append(lvl_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), font)
    fonts.set(qn("w:hAnsi"), font)
    r_pr.append(fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_id = OxmlElement("w:abstractNumId")
    abs_id.set(qn("w:val"), str(abstract_id))
    num.append(abs_id)
    numbering.append(num)


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])


def add_para(doc, text="", bold_lead=None, align=None, after=6, italic=False, keep_next=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.keep_with_next = keep_next
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic)
    return p


def add_list(doc, items, kind="bullet"):
    num_id = 101 if kind == "bullet" else 102 if kind == "decimal" else 103
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        p.paragraph_format.keep_together = True
        apply_num(p, num_id)
        r = p.add_run(item)
        set_run_font(r)


def add_chapter(doc, roman, title):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = False
    p.paragraph_format.keep_with_next = True
    label = f"CAPÍTULO {roman}\n{title.upper()}" if roman else title.upper()
    r = p.add_run(label)
    set_run_font(r, size=16, color=GREEN, bold=True)
    return p


def add_article(doc, number, title, paragraphs=None, bullets=None, decimals=None, letters=None):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"Artículo {number}. ({title})")
    set_run_font(r, size=12.5, color=NAVY, bold=True)
    for text in paragraphs or []:
        add_para(doc, text)
    if bullets:
        add_list(doc, bullets, "bullet")
    if decimals:
        add_list(doc, decimals, "decimal")
    if letters:
        add_list(doc, letters, "letter")


def add_table(doc, headers, rows, widths_dxa, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, header in enumerate(headers):
        set_cell_shading(hdr.cells[i], GREEN)
        p = hdr.cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=font_size, color=WHITE, bold=True)
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        if row_idx % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, PALE)
        for i, value in enumerate(values):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    set_table_geometry(table, widths_dxa)
    add_para(doc, "", after=2)
    return table


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GREEN)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(label + ": ")
    set_run_font(r1, color=GREEN, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, color=NAVY)
    add_para(doc, "", after=2)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    h1 = styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor.from_string(GREEN)
    h1._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    h1._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    h2.font.name = "Calibri"
    h2.font.size = Pt(12.5)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor.from_string(NAVY)
    h2._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    h2._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    h3 = styles["Heading 3"]
    h3.font.name = "Calibri"
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor.from_string(NAVY)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True


def configure_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("INSO | REGLAMENTO DE USO DEL SISTEMA CERTIFICADOR")
    set_run_font(r, size=8.5, color=GRAY, bold=True)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), MID_GREEN)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer = section.footer
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(table, [6800, 2560], indent_dxa=0)
    set_repeat_table_header(table.rows[0])
    set_table_borders(table, color=WHITE, size="0")
    left = table.cell(0, 0).paragraphs[0]
    left.paragraph_format.space_after = Pt(0)
    r = left.add_run("Código: INSO-DTI-RGL-001 | Versión 0 | Estado: Proyecto")
    set_run_font(r, size=8, color=GRAY)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.paragraph_format.space_after = Pt(0)
    r = right.add_run("Página ")
    set_run_font(r, size=8, color=GRAY)
    f = add_field(right, "PAGE")
    set_run_font(f, size=8, color=GRAY)


def build():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    configure_section(section)
    configure_header_footer(section)
    add_num_definition(doc, 101, 201, "bullet", "•", left=720, hanging=360)
    add_num_definition(doc, 102, 202, "decimal", "%1.", left=720, hanging=360)
    add_num_definition(doc, 103, 203, "lowerLetter", "%1)", left=720, hanging=360)

    # Portada editorial.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if LOGO.exists():
        picture = p.add_run().add_picture(str(LOGO), width=Inches(2.05))
        doc_pr = picture._inline.docPr
        doc_pr.set("title", "Logotipo del INSO")
        doc_pr.set("descr", "Logotipo institucional del Instituto Nacional de Salud Ocupacional")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("INSTITUTO NACIONAL DE SALUD OCUPACIONAL")
    set_run_font(r, size=12, color=GREEN, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("REGLAMENTO DE USO DEL\nSISTEMA CERTIFICADOR INSO")
    set_run_font(r, size=25, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run("Proyecto normativo adaptado a la operación real del sistema")
    set_run_font(r, size=12.5, color=GRAY, italic=True)

    cover_table = doc.add_table(rows=4, cols=2)
    set_table_geometry(cover_table, [2700, 6660])
    set_repeat_table_header(cover_table.rows[0])
    for idx, (label, value) in enumerate([
        ("Código", "INSO-DTI-RGL-001"),
        ("Versión", "0 - Proyecto para revisión"),
        ("Fecha de elaboración", "26 de agosto de 2026"),
        ("Aprobación", "Pendiente de informe técnico, informe legal y resolución administrativa"),
    ]):
        set_cell_shading(cover_table.cell(idx, 0), LIGHT_GREEN)
        p1 = cover_table.cell(idx, 0).paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        set_run_font(p1.add_run(label), size=9.5, color=GREEN, bold=True)
        p2 = cover_table.cell(idx, 1).paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        set_run_font(p2.add_run(value), size=9.5)

    add_para(doc, "La Paz - Bolivia", align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    doc.add_page_break()

    # Control y nota de uso.
    add_chapter(doc, "", "CONTROL DEL DOCUMENTO")
    add_callout(
        doc,
        "Naturaleza",
        "El presente texto es un proyecto de reglamento. Debe ser revisado por las áreas técnica, de planificación y jurídica del INSO antes de su aprobación y entrada en vigencia.",
    )
    add_table(
        doc,
        ["Etapa", "Responsable", "Nombre y cargo", "Fecha / firma"],
        [
            ("Elaboración técnica", "Unidad o Área de Sistemas", "[Completar]", "[Completar]"),
            ("Revisión funcional", "Unidad responsable de certificación", "[Completar]", "[Completar]"),
            ("Revisión jurídica", "Asesoría Jurídica", "[Completar]", "[Completar]"),
            ("Aprobación", "Dirección General Ejecutiva", "[Completar]", "[Completar]"),
        ],
        [1800, 2600, 3000, 1960],
        font_size=8.8,
    )

    p = doc.add_paragraph(style="Heading 1")
    set_run_font(p.add_run("ÍNDICE DE CAPÍTULOS"), size=16, color=GREEN, bold=True)
    add_list(doc, [
        "Capítulo I. Disposiciones generales",
        "Capítulo II. Gobierno y responsabilidades del sistema",
        "Capítulo III. Usuarios, registro, roles y control de acceso",
        "Capítulo IV. Reglas de uso, obligaciones y prohibiciones",
        "Capítulo V. Gestión de trámites, requisitos y evidencias",
        "Capítulo VI. Revisión, pagos y emisión de certificados",
        "Capítulo VII. Información, confidencialidad y trazabilidad",
        "Capítulo VIII. Operación, soporte, continuidad e incidentes",
        "Capítulo IX. Control, responsabilidad y medidas por incumplimiento",
        "Capítulo X. Disposiciones transitorias y finales",
        "Anexos I a III",
    ], "decimal")

    p = doc.add_paragraph(style="Heading 1")
    set_run_font(p.add_run("NOTA DE FUNDAMENTACIÓN Y ADAPTACIÓN"), size=16, color=GREEN, bold=True)
    add_para(doc, "Este reglamento fue preparado a partir del análisis funcional del Sistema Certificador INSO y de modelos públicos de otras instituciones bolivianas. La adaptación prioriza controles compatibles con la arquitectura observada: cuentas individuales, roles y permisos, solicitantes y tramitadores autorizados, funcionarios, expedientes, requisitos, evidencias, pagos, revisión técnica, observaciones, historial, notificaciones y emisión de certificados.")
    add_para(doc, "No se realizó una transcripción literal de los modelos consultados. Se incorporaron principios comunes y se adecuaron a las funciones del INSO, manteniendo una separación clara entre regulación de uso, procedimiento sustantivo de certificación y política técnica de seguridad.")
    add_table(
        doc,
        ["Modelo institucional", "Elemento tomado como referencia", "Adaptación al Sistema Certificador INSO"],
        [
            ("DGAC - Reglamento de uso de sistemas informáticos", "Cuenta personal, acceso autorizado, responsabilidades y sanciones.", "Cuentas individuales; prohibición de compartir credenciales; ciclo de alta, modificación y baja."),
            ("Ministerio de Educación - Portal de Trámites", "Responsabilidad del usuario y del área de sistemas; roles; expediente de inicio a conclusión.", "Responsables funcionales, técnicos y usuarios; flujo completo de certificación."),
            ("AGETIC - Plataforma gob.bo", "Gestión institucional de usuarios, roles y actualización de información.", "Asignación documentada de roles y revisión periódica de permisos."),
            ("CTIC-EPB - Lineamientos de seguridad", "Credenciales, mínimo privilegio, respaldos, continuidad y gestión de incidentes.", "Controles proporcionales, reporte de incidentes y recuperación documentada."),
        ],
        [2200, 3300, 3860],
        font_size=8.5,
    )

    # CAPÍTULO I
    add_chapter(doc, "I", "DISPOSICIONES GENERALES")
    add_article(doc, 1, "Objeto", ["El presente Reglamento tiene por objeto regular el acceso, uso, administración, operación, control y seguridad del Sistema Certificador del Instituto Nacional de Salud Ocupacional - INSO, en adelante el Sistema, así como establecer los derechos, obligaciones y responsabilidades de las personas usuarias y de las unidades institucionales que intervienen en su gestión."])
    add_article(doc, 2, "Finalidad", ["El Reglamento tiene por finalidad asegurar que los trámites y certificados gestionados mediante el Sistema se procesen con legalidad, integridad, trazabilidad, oportunidad, confidencialidad, disponibilidad, responsabilidad y orientación a la ciudadanía."])
    add_article(doc, 3, "Ámbito de aplicación", ["El Reglamento es de cumplimiento obligatorio para servidoras y servidores públicos, personal eventual, consultores, administradores técnicos, solicitantes, beneficiarios, representantes legales, tramitadores autorizados y cualquier otra persona natural o jurídica que acceda al Sistema o intervenga en sus procesos."])
    add_article(doc, 4, "Marco normativo", ["El presente Reglamento se interpreta y aplica en concordancia con la Constitución Política del Estado; la Ley N° 1178 de Administración y Control Gubernamentales; la Ley N° 164 General de Telecomunicaciones, Tecnologías de Información y Comunicación; la Ley N° 1080 de Ciudadanía Digital; la Ley N° 2341 de Procedimiento Administrativo y su reglamentación, cuando corresponda; el Decreto Supremo N° 1793; el Decreto Supremo N° 2514 y sus modificaciones; el Decreto Supremo N° 3525; el Decreto Supremo N° 26119 relativo al INSO; y demás normativa vigente aplicable a gobierno electrónico, seguridad de la información, archivo, transparencia, responsabilidad por la función pública y certificación institucional.", "La unidad jurídica deberá verificar y actualizar las referencias normativas antes de la aprobación del Reglamento."])
    add_article(doc, 5, "Principios", bullets=[
        "Legalidad y competencia: cada actuación será realizada por la autoridad o persona habilitada dentro de sus atribuciones.",
        "Responsabilidad: toda cuenta y acción será atribuible a su titular o al responsable que la ejecute.",
        "Mínimo privilegio: los accesos se limitarán a lo estrictamente necesario para la función asignada.",
        "Segregación de funciones: se evitará que una sola persona concentre actividades incompatibles de registro, revisión, validación, emisión y administración técnica.",
        "Integridad y trazabilidad: se preservará la exactitud, secuencia y evidencia de las actuaciones.",
        "Confidencialidad y privacidad: la información se tratará según su naturaleza y necesidad institucional.",
        "Disponibilidad y continuidad: se adoptarán medidas razonables para mantener y recuperar el servicio.",
        "Buena fe y veracidad: las personas usuarias responden por la autenticidad y suficiencia de la información presentada.",
        "Accesibilidad y orientación a la ciudadanía: la gestión digital procurará claridad, facilidad de uso y atención no discriminatoria.",
    ])
    add_article(doc, 6, "Definiciones", paragraphs=[
        "Administrador funcional. Persona designada para configurar catálogos, requisitos, tipos de certificado, plantillas, flujos, roles o parámetros de negocio, conforme a autorización expresa.",
        "Administrador técnico. Personal responsable de la operación tecnológica, soporte, seguridad, mantenimiento, respaldos y recuperación del Sistema, sin competencia para resolver el fondo de un trámite.",
        "Certificado. Documento institucional emitido por autoridad competente como resultado de un trámite aprobado y registrado en el Sistema.",
        "Credencial. Conjunto de elementos de autenticación vinculados a una cuenta, tales como nombre de usuario, contraseña y, cuando esté habilitado, segundo factor.",
        "Evidencia. Archivo, dato, documento o registro presentado o generado para acreditar el cumplimiento de un requisito.",
        "Expediente electrónico. Conjunto ordenado de datos, documentos, evidencias, actuaciones, observaciones, pagos, notificaciones e historial de un trámite.",
        "Incidente de seguridad. Evento que compromete o puede comprometer la confidencialidad, integridad, disponibilidad o autenticidad de la información o del Sistema.",
        "Rol. Conjunto de permisos que habilita acciones determinadas dentro del Sistema.",
        "Solicitante. Persona natural o jurídica que inicia un trámite por cuenta propia o mediante representación válida.",
        "Tramitador. Persona natural autorizada por una empresa o solicitante para gestionar trámites dentro de los límites de su autorización.",
    ])
    add_article(doc, 7, "Carácter complementario", ["Este Reglamento regula el uso de la herramienta tecnológica. Los requisitos, plazos, aranceles, criterios técnicos, competencias y efectos jurídicos de cada certificación se rigen por la normativa y los procedimientos específicos aprobados para cada servicio. En caso de contradicción, prevalecerá la norma de mayor jerarquía."])

    # CAPÍTULO II
    add_chapter(doc, "II", "GOBIERNO Y RESPONSABILIDADES DEL SISTEMA")
    add_article(doc, 8, "Autoridad institucional", ["La Dirección General Ejecutiva del INSO es responsable de aprobar el presente Reglamento, disponer su implementación y asegurar la coordinación entre las áreas técnicas, funcionales, jurídicas, administrativas y de seguridad de la información."])
    add_article(doc, 9, "Responsable funcional", bullets=[
        "Definir y mantener los flujos, estados, requisitos, responsables y reglas de cada tipo de certificación.",
        "Validar los cambios funcionales antes de su puesta en producción.",
        "Solicitar y justificar roles y permisos del personal a su cargo.",
        "Supervisar la calidad, oportunidad y consistencia de la información del trámite.",
        "Coordinar la capacitación y comunicar cambios de procedimiento.",
        "Resolver consultas funcionales sin invadir competencias jurídicas o técnicas especializadas.",
    ])
    add_article(doc, 10, "Responsable técnico o Área de Sistemas", bullets=[
        "Administrar la infraestructura, aplicaciones, base de datos, configuración técnica y mecanismos de autenticación.",
        "Ejecutar altas, modificaciones, bloqueos y bajas de acceso debidamente autorizadas.",
        "Aplicar actualizaciones, correcciones, medidas de protección y monitoreo.",
        "Gestionar respaldos, pruebas de restauración, continuidad y recuperación.",
        "Registrar, atender y escalar incidentes de seguridad y solicitudes de soporte.",
        "Conservar registros técnicos de administración y evitar modificaciones directas de datos sin autorización, justificación y trazabilidad.",
        "Mantener separación entre privilegios técnicos y decisiones sustantivas de certificación.",
    ])
    add_article(doc, 11, "Responsable de seguridad de la información", ["La persona designada como Responsable de Seguridad de la Información coordinará la gestión de riesgos, controles, incidentes, concientización y cumplimiento del Plan Institucional de Seguridad de la Información. Los incidentes sujetos a reporte externo serán comunicados al Centro de Gestión de Incidentes Informáticos dentro del plazo y por los canales previstos en la normativa vigente."])
    add_article(doc, 12, "Jefaturas y responsables de unidad", bullets=[
        "Solicitar accesos únicamente para personal que lo requiera.",
        "Revisar periódicamente la vigencia de funciones, roles y permisos.",
        "Comunicar de inmediato vacaciones prolongadas, cambios de cargo, desvinculaciones, suspensiones o cualquier situación que requiera ajustar accesos.",
        "Supervisar el uso correcto del Sistema en su unidad y reportar irregularidades.",
        "Asegurar la transferencia formal de trámites pendientes cuando exista cambio de responsable.",
    ])

    # CAPÍTULO III
    add_chapter(doc, "III", "USUARIOS, REGISTRO, ROLES Y CONTROL DE ACCESO")
    add_article(doc, 13, "Categorías de usuario", ["El Sistema podrá habilitar, según su configuración vigente, las categorías de Super Administrador, Administrador, Funcionario, Solicitante y Tramitador, además de permisos específicos asociados a consulta, inicio, atención, gestión, historial, validación de pagos, validación de tramitadores, emisión de certificados, reportes y configuración."])
    add_article(doc, 14, "Creación de cuentas", ["Toda cuenta deberá corresponder a una persona identificable y a un correo electrónico válido. Las cuentas internas requerirán solicitud y autorización de la jefatura competente. Las cuentas externas deberán cumplir los mecanismos de registro, verificación y acreditación previstos por el INSO.", "Queda prohibida la creación de cuentas genéricas o compartidas, salvo cuentas técnicas no interactivas expresamente autorizadas, documentadas y controladas por el Área de Sistemas."])
    add_article(doc, 15, "Asignación de roles y permisos", ["La asignación de roles y permisos se basará en funciones, competencia y necesidad de acceso. Toda ampliación de privilegios deberá contar con justificación y autorización documentada. Los permisos directos a una cuenta serán excepcionales y estarán sujetos a revisión reforzada."])
    add_article(doc, 16, "Representación y tramitadores", ["El acceso de un tramitador a trámites de una empresa requerirá autorización vigente y verificable. La habilitación no transfiere al tramitador la titularidad del trámite ni exime al solicitante o representante legal de sus responsabilidades.", "La empresa o persona representada deberá comunicar la revocatoria, suspensión o modificación de la autorización. El INSO podrá requerir confirmación adicional cuando existan inconsistencias, riesgo de suplantación o cambio de representantes."])
    add_article(doc, 17, "Responsabilidad sobre credenciales", bullets=[
        "La credencial es personal, individual e intransferible.",
        "La persona usuaria deberá mantenerla en secreto, evitar su almacenamiento visible y utilizar mecanismos seguros de recuperación.",
        "No podrá comunicar contraseñas por correo no protegido, mensajería informal, teléfono o documentos de libre acceso.",
        "Cuando exista segundo factor de autenticación, deberá custodiar también el dispositivo, código o medio asociado.",
        "Toda actividad realizada después de una autenticación válida se presumirá efectuada por el titular, salvo prueba de compromiso reportado oportunamente.",
    ])
    add_article(doc, 18, "Contraseñas y autenticación", ["Las contraseñas deberán cumplir la longitud, complejidad, historial y demás controles definidos por la política institucional y la configuración vigente del Sistema. El INSO podrá implementar CAPTCHA, limitación de intentos, verificación de correo, segundo factor, cierre de sesiones y otros controles proporcionales al riesgo.", "Las contraseñas temporales deberán cambiarse en el primer acceso o dentro del plazo establecido. Ante sospecha de compromiso, la persona usuaria deberá cambiarla inmediatamente y reportar el hecho."])
    add_article(doc, 19, "Revisión, suspensión y baja", ["Los derechos de acceso serán revisados de forma periódica y cuando cambien las funciones o la relación con el INSO. La cuenta podrá ser suspendida preventivamente por inactividad prolongada, intentos anómalos, riesgo de seguridad, documentación de representación vencida, incumplimiento o solicitud autorizada.", "La baja de personal interno deberá gestionarse de inmediato. Antes de la desactivación se asegurará la reasignación de trámites, la preservación de registros y la revocación de sesiones o mecanismos de acceso vigentes."])
    add_article(doc, 20, "Sesiones y dispositivos", bullets=[
        "Cerrar la sesión al concluir el trabajo o al dejar el dispositivo sin supervisión.",
        "No guardar credenciales en equipos de uso público o compartido.",
        "Mantener el navegador y sistema operativo razonablemente actualizados y libres de software malicioso.",
        "Evitar redes inseguras para transmitir documentación sensible; cuando no exista alternativa, utilizar los mecanismos de protección dispuestos por el INSO.",
        "No eludir controles de sesión, autenticación o restricción geográfica o técnica.",
    ])

    # CAPÍTULO IV
    add_chapter(doc, "IV", "REGLAS DE USO, OBLIGACIONES Y PROHIBICIONES")
    add_article(doc, 21, "Uso autorizado", ["El Sistema será utilizado exclusivamente para las finalidades institucionales y los trámites habilitados. Las personas usuarias internas actuarán dentro de sus funciones; las externas, dentro del alcance de su solicitud o representación."])
    add_article(doc, 22, "Obligaciones generales", bullets=[
        "Proporcionar información completa, actual, legible, pertinente y veraz.",
        "Revisar los datos antes de guardar o enviar una actuación.",
        "Atender observaciones y notificaciones dentro de los plazos aplicables.",
        "Conservar los originales cuando la normativa o el INSO pueda requerir su cotejo.",
        "Proteger información personal, técnica, comercial o institucional a la que se tenga acceso.",
        "Reportar errores, accesos indebidos, pérdida de credenciales o comportamientos anómalos.",
        "Utilizar los canales de soporte establecidos y colaborar con verificaciones legítimas.",
    ])
    add_article(doc, 23, "Prohibiciones", bullets=[
        "Compartir, prestar, vender o ceder cuentas, credenciales, códigos o sesiones.",
        "Suplantar a otra persona o actuar sin representación válida.",
        "Alterar, fabricar, ocultar o eliminar información o evidencias para inducir a error.",
        "Intentar acceder a trámites, módulos, archivos, reportes o datos no autorizados.",
        "Probar vulnerabilidades, interceptar comunicaciones, automatizar consultas o sobrecargar el servicio sin autorización escrita.",
        "Introducir código malicioso, archivos dañinos o contenido ajeno al trámite.",
        "Modificar registros mediante acceso directo a la base de datos o herramientas externas sin procedimiento autorizado.",
        "Utilizar datos del Sistema con fines particulares, comerciales, políticos, discriminatorios o contrarios al ordenamiento jurídico.",
        "Extraer masivamente información o reproducir documentos internos sin autorización.",
        "Desactivar, eludir o interferir controles de seguridad, auditoría o trazabilidad.",
    ])
    add_article(doc, 24, "Uso de información obtenida", ["La información a la que se acceda por razón del cargo, trámite o representación solo podrá usarse para la finalidad que justificó el acceso. La obligación de reserva subsistirá después de la conclusión del trámite, baja de cuenta, revocatoria de representación o desvinculación laboral."])
    add_article(doc, 25, "Propiedad y custodia institucional", ["El software, configuraciones, catálogos, plantillas, registros, reportes y documentación producida por el personal en ejercicio de funciones constituyen activos institucionales, sin perjuicio de los derechos de las personas titulares sobre sus datos y documentos conforme al ordenamiento vigente."])

    # CAPÍTULO V
    add_chapter(doc, "V", "GESTIÓN DE TRÁMITES, REQUISITOS Y EVIDENCIAS")
    add_article(doc, 26, "Inicio del trámite", ["El trámite se considerará iniciado cuando la solicitud haya sido registrada y enviada conforme al flujo habilitado, se hayan completado los datos mínimos y se hayan adjuntado los requisitos exigibles. Los borradores no enviados carecen de efecto de presentación, salvo disposición expresa del procedimiento específico."])
    add_article(doc, 27, "Identificación del beneficiario y la empresa", ["La persona usuaria deberá seleccionar o registrar correctamente al beneficiario, empresa, fabricante, producto, procedencia y demás datos aplicables. El Sistema podrá impedir el inicio cuando no exista una relación activa de representación o cuando la información esencial sea inconsistente."])
    add_article(doc, 28, "Documentos y evidencias", bullets=[
        "Los archivos deberán corresponder al requisito indicado y ser legibles, completos y libres de alteraciones fraudulentas.",
        "La carga de un archivo no sustituye las formalidades de firma, legalización, vigencia, traducción o cotejo que exija la normativa específica.",
        "La persona usuaria es responsable de verificar que el archivo correcto haya sido incorporado antes del envío.",
        "El personal revisor deberá evaluar la evidencia según el alcance de su competencia y registrar el resultado y las observaciones de manera clara.",
        "No se cargarán datos personales o sensibles que no sean necesarios para el trámite.",
    ])
    add_article(doc, 29, "Integridad del expediente", ["Las actuaciones del trámite formarán un expediente electrónico ordenado. No podrán eliminarse u ocultarse actuaciones válidamente registradas para alterar su secuencia. Las correcciones deberán realizarse mediante funciones autorizadas que conserven la trazabilidad del cambio."])
    add_article(doc, 30, "Observaciones y subsanación", ["Las observaciones deberán ser específicas, comprensibles, relacionadas con un requisito o dato y emitidas por persona competente. Cuando el flujo lo permita, el personal revisor registrará primero el resultado técnico y luego notificará formalmente la necesidad de corrección.", "La persona solicitante o tramitadora deberá subsanar únicamente los aspectos habilitados, sin reemplazar evidencia aprobada o alterar información no observada salvo autorización."])
    add_article(doc, 31, "Derivación y asignación", ["La asignación o derivación a personal técnico o administrativo deberá responder a competencia, carga de trabajo, área y ausencia de conflicto de interés. Toda reasignación relevante quedará registrada en el historial del trámite."])
    add_article(doc, 32, "Historial y seguimiento", ["El historial permitirá conocer los estados, responsables y principales actuaciones del trámite. Los registros de fecha y hora se generarán con la configuración oficial del Sistema y no podrán alterarse manualmente, salvo corrección técnica excepcional documentada."])

    # CAPÍTULO VI
    add_chapter(doc, "VI", "REVISIÓN, PAGOS Y EMISIÓN DE CERTIFICADOS")
    add_article(doc, 33, "Revisión técnica", ["La revisión técnica será realizada por personal autorizado con base en los requisitos, criterios y procedimientos vigentes. El Sistema es una herramienta de apoyo y trazabilidad; no reemplaza el juicio técnico ni permite apartarse de la normativa aplicable.", "La persona revisora deberá declarar impedimentos o conflictos de interés y abstenerse cuando corresponda."])
    add_article(doc, 34, "Validación de pagos", ["Cuando el trámite requiera pago, su registro y validación se realizarán por el rol autorizado y con respaldo verificable. La sola carga de un comprobante no equivale a validación ni genera por sí misma derecho a la emisión del certificado.", "Las diferencias, duplicidades, reversos o inconsistencias se tratarán mediante el procedimiento financiero y administrativo aplicable."])
    add_article(doc, 35, "Decisión y finalización", ["La aprobación, observación, rechazo o finalización de un trámite deberá ser adoptada por la autoridad o persona competente, con motivación suficiente cuando corresponda. El cambio de estado no deberá utilizarse para omitir revisiones, controles o autorizaciones requeridas."])
    add_article(doc, 36, "Emisión del certificado", ["El certificado se emitirá únicamente cuando el expediente cumpla los requisitos, se encuentren registradas las revisiones y validaciones exigibles y exista autorización conforme al procedimiento específico. La plantilla utilizada deberá corresponder al tipo de certificado vigente y estar controlada por el responsable funcional.", "La emisión deberá dejar registro de la persona responsable, fecha, trámite de origen y versión o datos necesarios para su identificación institucional."])
    add_article(doc, 37, "Corrección, anulación o reemisión", ["La corrección, anulación o reemisión de un certificado requerirá causa documentada, autorización competente y preservación del vínculo con el documento anterior. No procederá la sustitución informal de archivos ni la modificación directa de un certificado emitido.", "El procedimiento específico definirá los efectos frente a terceros, la comunicación a la persona interesada y, cuando corresponda, la recuperación o invalidación del documento anterior."])

    # CAPÍTULO VII
    add_chapter(doc, "VII", "INFORMACIÓN, CONFIDENCIALIDAD Y TRAZABILIDAD")
    add_article(doc, 38, "Clasificación y necesidad de acceso", ["La información será tratada de acuerdo con su carácter público, interno, confidencial o sensible. El acceso se concederá por necesidad funcional y se limitará al expediente, módulo o reporte necesario para cumplir la tarea asignada."])
    add_article(doc, 39, "Datos personales", ["El INSO y las personas usuarias tratarán datos personales únicamente para finalidades legítimas vinculadas al trámite, gestión institucional, control, seguridad o cumplimiento legal. Se evitará la recopilación excesiva y se aplicarán medidas para prevenir divulgación, pérdida, alteración o acceso no autorizado.", "Las solicitudes de acceso, rectificación u otras relacionadas con datos personales serán atendidas por el canal y autoridad competente, considerando las obligaciones de archivo, control y conservación."])
    add_article(doc, 40, "Confidencialidad", ["El personal con acceso a información técnica, comercial, personal o institucional no pública deberá guardar reserva y no podrá copiarla, fotografiarla, enviarla, imprimirla o revelarla salvo necesidad funcional y autorización."])
    add_article(doc, 41, "Registros de auditoría", ["El Sistema conservará, según sus capacidades, registros de creación, modificación, eliminación lógica, asignación, revisión, notificación, validación, emisión y administración. Los registros de auditoría serán protegidos contra alteración y consultados únicamente por personal autorizado para control, soporte, investigación o cumplimiento."])
    add_article(doc, 42, "Conservación y archivo", ["Los expedientes, certificados, respaldos y registros se conservarán durante los plazos definidos por la normativa archivística, sectorial, administrativa y de control aplicable. La eliminación o depuración deberá estar autorizada, documentada y ejecutada de forma segura."])
    add_article(doc, 43, "Intercambio y reportes", ["La exportación, intercambio o remisión de información a otras entidades requerirá competencia, finalidad definida, datos mínimos necesarios y un canal autorizado. Los reportes institucionales deberán respetar las restricciones de acceso y, cuando sea posible, aplicar agregación o anonimización."])

    # CAPÍTULO VIII
    add_chapter(doc, "VIII", "OPERACIÓN, SOPORTE, CONTINUIDAD E INCIDENTES")
    add_article(doc, 44, "Disponibilidad y mantenimiento", ["El INSO procurará la disponibilidad y continuidad razonable del Sistema. Podrá realizar mantenimientos, actualizaciones o suspensiones programadas, comunicándolos por los medios disponibles cuando afecten la atención. La indisponibilidad técnica no implica aprobación automática ni altera plazos salvo decisión de autoridad competente."])
    add_article(doc, 45, "Soporte", ["Las consultas e incidentes deberán registrarse por los canales oficiales de soporte, incluyendo identificación de la persona usuaria, módulo afectado, fecha, descripción y evidencia mínima necesaria. No deberán enviarse contraseñas ni información sensible innecesaria.", "El soporte técnico no podrá modificar decisiones de trámite ni efectuar actuaciones en nombre de la persona usuaria, excepto procedimiento de asistencia autorizado y trazable."])
    add_article(doc, 46, "Respaldos y restauración", bullets=[
        "Identificar información y configuraciones críticas.",
        "Definir periodicidad, retención y ubicación de copias conforme al riesgo y obligaciones institucionales.",
        "Proteger los respaldos contra acceso no autorizado, alteración y pérdida.",
        "Realizar pruebas periódicas de restauración y documentar sus resultados.",
        "Mantener copias separadas o redundantes cuando la criticidad lo justifique.",
    ])
    add_article(doc, 47, "Gestión de incidentes", ["Toda persona usuaria deberá reportar inmediatamente pérdida de credenciales, acceso ajeno, archivos sospechosos, divulgación, alteración, indisponibilidad anormal o cualquier indicio de incidente. Deberá preservar la evidencia disponible y abstenerse de investigar por cuenta propia si ello puede agravar el hecho.", "El responsable competente registrará, clasificará, contendrá, analizará, erradicará y recuperará el servicio; documentará las acciones y comunicará el incidente conforme a la normativa. Cuando corresponda, el reporte al CGII se realizará en el plazo vigente."])
    add_article(doc, 48, "Continuidad y contingencia", ["El INSO contará con procedimientos para mantener o restablecer procesos críticos de certificación ante fallas tecnológicas, desastres, indisponibilidad de proveedores o incidentes. Los procedimientos manuales o alternos deberán definir autorización, custodia, posterior regularización en el Sistema y controles para evitar duplicidad."])

    # CAPÍTULO IX
    add_chapter(doc, "IX", "CONTROL, RESPONSABILIDAD Y MEDIDAS POR INCUMPLIMIENTO")
    add_article(doc, 49, "Supervisión y monitoreo", ["El INSO podrá monitorear accesos, eventos y uso del Sistema con fines de seguridad, control, continuidad, auditoría y cumplimiento, respetando la finalidad institucional y la normativa aplicable. Las personas usuarias no deberán esperar privacidad sobre actividades realizadas dentro del Sistema, sin perjuicio de la protección de sus datos personales."])
    add_article(doc, 50, "Verificación e investigación", ["Ante indicios de incumplimiento, la autoridad competente podrá preservar registros, bloquear preventivamente accesos, solicitar informes y realizar verificaciones. Las medidas deberán ser proporcionales, documentadas y respetar el derecho a la defensa y el debido proceso."])
    add_article(doc, 51, "Medidas aplicables a usuarios externos", bullets=[
        "Advertencia y requerimiento de regularización.",
        "Cierre de sesiones o cambio obligatorio de credenciales.",
        "Suspensión temporal de cuenta o de una representación.",
        "Revocatoria de acceso cuando desaparezca la causa de habilitación o se acredite uso indebido.",
        "Remisión de antecedentes a la autoridad competente cuando exista posible responsabilidad civil, administrativa o penal.",
    ])
    add_article(doc, 52, "Responsabilidad del personal interno", ["El incumplimiento por servidoras y servidores públicos, personal eventual o consultores será tratado conforme al Reglamento Interno de Personal, régimen de responsabilidad por la función pública, contratos y demás normativa vigente, sin perjuicio de otras acciones legales."])
    add_article(doc, 53, "No alteración del trámite", ["La suspensión de una cuenta no anula automáticamente los actos válidamente realizados ni extingue obligaciones, plazos o responsabilidades. La autoridad funcional determinará la continuidad, reasignación, observación o revisión del trámite según el caso."])

    # CAPÍTULO X
    add_chapter(doc, "X", "DISPOSICIONES TRANSITORIAS Y FINALES")
    add_article(doc, 54, "Aceptación y conocimiento", ["El acceso al Sistema implica conocimiento y obligación de cumplimiento del presente Reglamento y de los avisos o términos específicos que se publiquen. El INSO podrá requerir constancia expresa de aceptación al crear o renovar una cuenta."])
    add_article(doc, 55, "Capacitación", ["Las áreas responsables organizarán actividades de inducción y actualización para usuarios internos y materiales de orientación para usuarios externos. Los roles críticos o privilegiados requerirán capacitación específica antes o inmediatamente después de su habilitación."])
    add_article(doc, 56, "Revisión y actualización", ["El Reglamento será revisado al menos cada dos años o antes cuando existan cambios normativos, funcionales, tecnológicos, organizacionales, incidentes relevantes o resultados de auditoría que lo justifiquen. Toda versión aprobada deberá identificar fecha, autoridad y cambios principales."])
    add_article(doc, 57, "Interpretación", ["Las dudas sobre aspectos funcionales serán atendidas por la unidad responsable del proceso; las técnicas, por el Área de Sistemas; las de seguridad, por el Responsable de Seguridad de la Información; y las jurídicas, por Asesoría Jurídica. Ninguna interpretación podrá ampliar privilegios sin autorización formal."])
    add_article(doc, 58, "Vigencia", ["El presente Reglamento entrará en vigencia a partir de su aprobación mediante la resolución correspondiente y su comunicación o publicación por los medios institucionales."])

    p = doc.add_paragraph(style="Heading 2")
    set_run_font(p.add_run("DISPOSICIÓN TRANSITORIA ÚNICA"), size=12.5, color=NAVY, bold=True)
    add_para(doc, "Dentro de los noventa (90) días calendario posteriores a la aprobación, las áreas responsables deberán: actualizar el inventario de cuentas y privilegios; regularizar autorizaciones de tramitadores; identificar roles críticos; definir canales de soporte e incidentes; aprobar el procedimiento de altas, modificaciones y bajas; y ejecutar una prueba documentada de restauración de respaldos. El plazo podrá ajustarse mediante resolución o instrucción expresa fundada.")

    # ANEXO I
    add_chapter(doc, "", "ANEXO I. MATRIZ MÍNIMA DE ROLES Y RESPONSABILIDADES")
    add_para(doc, "La matriz deberá alinearse con la configuración vigente del Sistema y con la estructura organizacional aprobada. No sustituye la autorización individual de accesos.")
    add_table(
        doc,
        ["Rol", "Alcance esperado", "Controles mínimos"],
        [
            ("Super Administrador", "Acceso excepcional y total para configuración crítica y recuperación.", "Designación expresa; uso restringido; segundo factor cuando esté disponible; revisión frecuente; registro reforzado."),
            ("Administrador", "Usuarios, roles, permisos, catálogos y configuración autorizada.", "Separación del trámite sustantivo; autorización de cambios; revisión periódica."),
            ("Funcionario", "Consulta, registro, revisión, derivación o gestión según cargo y permisos.", "Mínimo privilegio; bandeja individual; conflicto de interés; supervisión de jefatura."),
            ("Solicitante", "Inicio y seguimiento de trámites propios o de empresa representada.", "Identidad verificada; correo válido; relación empresarial activa; responsabilidad por datos y evidencias."),
            ("Tramitador", "Gestión por empresas que lo autorizaron.", "Autorización verificable, alcance definido, vigencia y revocatoria; sin transferencia de titularidad."),
            ("Validador de pagos", "Registro y validación de pagos o comprobantes.", "Separación respecto de la emisión cuando sea viable; respaldo y conciliación."),
            ("Emisor de certificados", "Emisión y envío de certificados aprobados.", "Competencia expresa; expediente completo; plantilla vigente; trazabilidad."),
        ],
        [2100, 3500, 3760],
        font_size=8.4,
    )

    # ANEXO II
    add_chapter(doc, "", "ANEXO II. CONSTANCIA DE CONOCIMIENTO Y USO RESPONSABLE")
    add_para(doc, "Yo, ________________________________________________, con documento de identidad N° ____________________, en mi condición de ________________________________________________, declaro haber recibido, leído y comprendido el Reglamento de Uso del Sistema Certificador INSO.")
    add_para(doc, "Me comprometo a utilizar mi cuenta de forma personal, custodiar mis credenciales, actuar dentro de los permisos asignados, proteger la información y reportar incidentes o accesos indebidos. Reconozco que las acciones realizadas con mi cuenta quedan sujetas a trazabilidad y a las responsabilidades previstas por la normativa vigente.")
    add_para(doc, "Usuario o correo institucional: ________________________________________________")
    add_para(doc, "Rol(es) autorizado(s): ________________________________________________________")
    add_para(doc, "Lugar y fecha: ________________________________________________________________")
    add_para(doc, "Firma de la persona usuaria: _____________________  Firma de autorización: _____________________")

    # ANEXO III
    add_chapter(doc, "", "ANEXO III. FUENTES INSTITUCIONALES CONSULTADAS")
    add_para(doc, "Fuentes públicas utilizadas como modelos de estructura y controles. Consulta realizada el 26 de agosto de 2026. Las referencias deben verificarse nuevamente antes de la aprobación oficial.")
    sources = [
        ("Dirección General de Aeronáutica Civil (DGAC)", "Reglamento de Uso de Sistemas Informáticos, Servicios y Equipos de Computación", "https://www.dgac.gob.bo/wp-content/uploads/2022/02/DAFASIS-RGL-001R0-REGLAMENTO-DE-SISTEMAS-1.pdf"),
        ("Ministerio de Educación", "Manual de Usuario para el Portal de Trámites, aprobado mediante R.M. N° 0629/2025", "https://www.minedu.gob.bo/files/documentos-normativos/resoluciones-ministeriales/2025/RM_0629-2025_compressed.pdf"),
        ("AGETIC", "Lineamientos para el uso y funcionamiento de la Plataforma Digital del Estado Plurinacional de Bolivia gob.bo", "https://agetic.gob.bo/sites/default/files/2025-06/1.Lineamientos%20te%CC%81cnicos%20para%20el%20uso%20y%20funcionamiento%20de%20la%20plataforma_0.pdf"),
        ("CTIC-EPB / AGETIC", "Lineamientos para la elaboración e implementación de Planes Institucionales de Seguridad de la Información", "https://csirt.gob.bo/sites/default/files/2017-12/Documento_Seguridad_SEG-001.pdf"),
        ("AGETIC", "Términos y Condiciones de Uso de la Plataforma Digital gob.bo", "https://agetic.gob.bo/sites/default/files/2025-06/TE%CC%81RMINOS%20Y%20CONDICIONES%20DE%20USO%20DE%20LA%20PLATAFORMA%20DIGITAL%20%E2%80%9CGOB.BO%E2%80%9D-firmado.pdf"),
        ("Ministerio de Salud y Deportes", "Reglamento Interno de Uso y Administración de Recursos Informáticos", "https://www.minsalud.gob.bo/3045-reglamento-recursos-informaticos-pdf"),
        ("AGETIC", "Políticas, lineamientos y estándares: marco normativo de gobierno electrónico, Ciudadanía Digital y seguridad", "https://www.agetic.gob.bo/politicas-lineamientos-estandares/"),
        ("INSO / análisis técnico interno", "Rutas, roles, permisos, modelos auditables, flujos de trámite, notificaciones y emisión observados en el repositorio del Sistema Certificador", ""),
    ]
    for institution, title, url in sources:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.10
        r1 = p.add_run(institution + ". ")
        set_run_font(r1, bold=True)
        if url:
            add_hyperlink(p, title, url)
        else:
            r2 = p.add_run(title)
            set_run_font(r2, italic=True)
    add_callout(doc, "Criterio jurídico", "La consulta de modelos institucionales orienta la estructura y los controles, pero no reemplaza el análisis de competencia, legalidad y jerarquía normativa que debe realizar Asesoría Jurídica del INSO.")

    # Propiedades y guardado.
    props = doc.core_properties
    props.title = "Reglamento de Uso del Sistema Certificador INSO"
    props.subject = "Proyecto normativo de uso, administración y seguridad"
    props.author = "Instituto Nacional de Salud Ocupacional - INSO"
    props.keywords = "INSO, sistema certificador, reglamento, trámites, seguridad de la información"
    props.comments = "Proyecto preparado para revisión técnica y jurídica."

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
