from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("artifacts/Prohibiciones_y_Consecuencias_Usuarios_Externos_INSO.docx")

FONT = "Calibri"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "5B6573"
LIGHT = "F2F4F7"
PALE_BLUE = "E8EEF5"
WHITE = "FFFFFF"
CONTENT_DXA = 9360


def set_run_font(run, size=11, bold=None, italic=None, color=INK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_paragraph_callout(paragraph, fill=PALE_BLUE, border=DARK_BLUE):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)

    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = p_bdr.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        p_bdr.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), border)


def mark_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    r_pr.extend([r_fonts, color, underline, size])
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_run_font(run, size=9, color=MUTED)
    for field_name in ("PAGE", "NUMPAGES"):
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), field_name)
        r = OxmlElement("w:r")
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), FONT)
        r_fonts.set(qn("w:hAnsi"), FONT)
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "18")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), MUTED)
        r_pr.extend([r_fonts, size, color])
        r.append(r_pr)
        t = OxmlElement("w:t")
        t.text = "1"
        r.append(t)
        fld.append(r)
        paragraph._p.append(fld)
        if field_name == "PAGE":
            sep = paragraph.add_run(" de ")
            set_run_font(sep, size=9, color=MUTED)


def create_numbering(doc, num_format, level_text, left=720, hanging=360):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), num_format)
    txt = OxmlElement("w:lvlText")
    txt.set(qn("w:val"), level_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "100")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    lvl.extend([start, fmt, txt, suff, p_pr])
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_id = OxmlElement("w:abstractNumId")
    abs_id.set(qn("w:val"), str(abstract_id))
    num.append(abs_id)
    numbering.append(num)
    return num_id


def add_numbered_item(doc, text, num_id):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.167
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p._p.get_or_add_pPr().append(num_pr)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_article(doc, number, title, body=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"Artículo {number}. ({title}).")
    set_run_font(r, size=11, bold=True, color=DARK_BLUE)
    if body:
        r2 = p.add_run(f" {body}")
        set_run_font(r2)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_callout(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.10
    set_paragraph_callout(p)
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, bold=True, color=DARK_BLUE)
    r2 = p.add_run(text)
    set_run_font(r2)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def build_document():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("INSO | Sistema Certificador | Documento de trabajo")
    set_run_font(hr, size=8.5, color=MUTED)

    footer = section.footer
    add_page_field(footer.paragraphs[0])

    lower_num = create_numbering(doc, "lowerLetter", "%1)", left=720, hanging=360)
    decimal_num = create_numbering(doc, "decimal", "%1.", left=720, hanging=360)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("PROPUESTA DE CAPÍTULO REGLAMENTARIO")
    set_run_font(r, size=10, bold=True, color=BLUE)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(5)
    tr = title.add_run("PROHIBICIONES Y CONSECUENCIAS PARA USUARIOS EXTERNOS")
    set_run_font(tr, size=23, bold=True, color="000000")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    sr = subtitle.add_run("Reglamento de Uso del Sistema Certificador del INSO")
    set_run_font(sr, size=13, color=MUTED)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(14)
    mr1 = meta.add_run("Estado: ")
    set_run_font(mr1, size=10, bold=True, color=MUTED)
    mr2 = meta.add_run("Texto propuesto para revisión y aprobación institucional")
    set_run_font(mr2, size=10, color=MUTED)

    add_callout(
        doc,
        "Criterio verificado",
        "La VUCE sí alcanza a personas naturales y jurídicas externas. Les exige presentar información correcta y asumir responsabilidad por su veracidad. No obstante, la VUCE no fija por sí misma penas concretas para cada incumplimiento: remite a las disposiciones legales vigentes.",
    )

    doc.add_heading("CAPÍTULO: USO INDEBIDO POR USUARIOS EXTERNOS", level=1)

    add_article(
        doc,
        "1",
        "Ámbito de aplicación",
        "El presente capítulo se aplica a solicitantes, representantes de empresas, tramitadores y demás personas externas que cuenten con acceso autorizado al Sistema Certificador.",
    )

    add_article(doc, "2", "Responsabilidades del usuario externo", "El usuario externo debe:")
    for item in (
        "Utilizar únicamente su propia cuenta y mantener sus credenciales bajo reserva.",
        "Presentar información y documentos completos, correctos, exactos, legibles y vigentes.",
        "Responder por la veracidad de la información y documentación incorporadas al trámite.",
        "Actuar únicamente por cuenta propia o de la empresa que le hubiera otorgado autorización vigente.",
        "Guardar confidencialidad respecto de la información a la que acceda mediante el sistema.",
    ):
        add_numbered_item(doc, item, lower_num)

    add_article(doc, "3", "Prohibiciones", "Queda prohibido al usuario externo:")
    for item in (
        "Compartir credenciales, permitir el uso de su cuenta o utilizar la identidad o cuenta de otra persona.",
        "Registrar información falsa o presentar documentos falsos, alterados, sustituidos o que no correspondan al trámite.",
        "Acceder o intentar acceder a cuentas, módulos, trámites, documentos o información para los cuales no tenga autorización.",
        "Modificar, suprimir, ocultar o inutilizar datos, documentos, antecedentes o registros del sistema.",
        "Iniciar o gestionar trámites en nombre de una persona o empresa sin autorización vigente.",
        "Copiar, descargar, utilizar o divulgar información de terceros para fines distintos al trámite autorizado.",
        "Manipular pagos, requisitos, estados del trámite o certificados; eludir controles de seguridad; introducir archivos dañinos; o ejecutar acciones destinadas a alterar o interrumpir el funcionamiento del sistema.",
    ):
        add_numbered_item(doc, item, lower_num)

    add_article(doc, "4", "Medidas de seguridad y consecuencias")
    add_body(
        doc,
        "I. Cuando existan indicios razonables de uso indebido, la unidad competente podrá cerrar sesiones, bloquear preventivamente la cuenta o los permisos afectados, resguardar los registros y revisar las operaciones relacionadas. Estas actuaciones son medidas temporales de seguridad y no constituyen una declaración de culpabilidad.",
    )
    add_body(doc, "II. Comprobado el incumplimiento, mediante el procedimiento aplicable, podrán disponerse:")
    for item in (
        "Amonestación escrita.",
        "Suspensión o desactivación de la cuenta o de los permisos de acceso.",
        "Baja de la autorización como tramitador, cuando corresponda.",
        "Remisión de antecedentes a la autoridad competente cuando existan indicios de responsabilidad civil, administrativa o penal.",
    ):
        add_numbered_item(doc, item, lower_num)
    add_body(
        doc,
        "III. La restricción de una cuenta no anula automáticamente un trámite o certificado. La revisión, rechazo, anulación o modificación de actuaciones deberá realizarse conforme al procedimiento específico que corresponda.",
    )

    add_article(doc, "5", "Debido proceso")
    add_body(
        doc,
        "Las consecuencias definitivas deberán ser dispuestas por autoridad competente, mediante decisión fundamentada y previa comunicación de los hechos al usuario, garantizando su derecho a presentar descargos. Las medidas urgentes de seguridad deberán ser temporales, proporcionales y registradas en el sistema.",
    )

    add_article(doc, "6", "Responsabilidad penal")
    add_body(
        doc,
        "Cuando los hechos puedan constituir delito, el INSO preservará los registros disponibles y remitirá los antecedentes al Ministerio Público u otra autoridad competente. El INSO no impone penas de cárcel; la responsabilidad y la pena solamente pueden ser determinadas dentro de un proceso penal.",
    )

    doc.add_heading("Alcance de las normas penales más relevantes", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    mark_table_header(table.rows[0])
    headers = ("Conducta", "Norma", "Consecuencia legal prevista")
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=9.5, bold=True, color=DARK_BLUE)

    rows = [
        (
            "Acceder, utilizar, modificar, suprimir o inutilizar datos sin autorización, causando perjuicio.",
            "Código Penal, art. 363 ter",
            "Prestación de trabajo hasta un año o multa hasta 200 días. Este artículo no establece cárcel.",
        ),
        (
            "Manipular datos para obtener beneficio indebido y provocar una transferencia patrimonial en perjuicio de tercero.",
            "Código Penal, art. 363 bis",
            "Reclusión de 1 a 5 años y multa de 60 a 200 días.",
        ),
        (
            "Falsificar o alterar documento público, o insertar declaraciones falsas en instrumento público.",
            "Código Penal, arts. 198 y 199",
            "Privación de libertad de 1 a 6 años, si concurren los elementos del delito.",
        ),
        (
            "Usar a sabiendas un documento falso o adulterado.",
            "Código Penal, art. 203",
            "La misma sanción correspondiente al autor de la falsedad aplicable.",
        ),
    ]
    for row_data in rows:
        row_cells = table.add_row().cells
        for idx, text in enumerate(row_data):
            p = row_cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(text)
            set_run_font(r, size=9.2)
    set_table_geometry(table, [3100, 2050, 4210])

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(8)
    note.paragraph_format.space_after = Pt(6)
    nr1 = note.add_run("Nota sobre el intento de acceso: ")
    set_run_font(nr1, bold=True, color=DARK_BLUE)
    nr2 = note.add_run(
        "un intento técnico puede justificar el bloqueo preventivo y la preservación de registros, pero no significa automáticamente que exista un delito con pena de cárcel. La tentativa penal exige los presupuestos del artículo 8 del Código Penal y debe ser valorada por las autoridades competentes."
    )
    set_run_font(nr2)

    doc.add_heading("Fundamento normativo verificado", level=1)

    add_numbered_item(
        doc,
        "Reglamento VUCE, versión 2, aprobado mediante Resolución de Directorio RD 01-062-25 de 29 de agosto de 2025: artículo 4 (incluye personas naturales y jurídicas), artículo 18 (exactitud y veracidad de documentos) y artículo 27 (confidencialidad y uso indebido).",
        decimal_num,
    )
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.space_after = Pt(6)
    add_hyperlink(p, "Fuente oficial de la Aduana Nacional", "https://www.aduana.gob.bo/sites/default/files/archivos/CIRCULARES%20221-2025_compressed_0.pdf")

    add_numbered_item(
        doc,
        "Ley N.º 2341 de Procedimiento Administrativo: artículos 72 a 76. Exige legalidad, tipicidad, proporcionalidad, presunción de inocencia y procedimiento previo para imponer sanciones administrativas.",
        decimal_num,
    )
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.space_after = Pt(6)
    add_hyperlink(p, "Fuente oficial de la Gaceta Oficial", "https://www.gacetaoficialdebolivia.gob.bo/normas/vergratis_gob/20285")

    add_numbered_item(
        doc,
        "Código Penal: artículos 8, 198, 199, 203, 363 bis y 363 ter. Su aplicación depende de los hechos probados y de la determinación de la autoridad penal competente.",
        decimal_num,
    )
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.space_after = Pt(6)
    add_hyperlink(p, "Texto oficial alojado por SEGIP", "https://www.segip.gob.bo/wp-content/uploads/2025/09/LEY-No-10426-DE-23-DE-AGOSTO-DE-1972-ELEVADO-A-RANGO-DE-LEY-Y-MODIFICADO-POR-LA-LEY-No-1768-DE-10-DE-MARZO-CODIGO-PENAL.pdf")

    add_numbered_item(
        doc,
        "Ley N.º 1080 de Ciudadanía Digital, artículos 9 y 11: responsabilidad sobre credenciales y sanción del uso indebido, suplantación o alteración. Esta referencia será directamente aplicable al acceso del sistema únicamente si se utiliza Ciudadanía Digital.",
        decimal_num,
    )
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.space_after = Pt(6)
    add_hyperlink(p, "Texto vigente en SILEP", "https://www.silep.gob.bo/norma/14001/ley_actualizada")

    add_callout(
        doc,
        "Advertencia de aprobación",
        "La VUCE se utiliza como modelo regulatorio; no otorga automáticamente al INSO competencia sancionadora. Antes de aprobar este capítulo, la Unidad Jurídica debe confirmar la competencia institucional y el procedimiento aplicable.",
    )

    doc.core_properties.title = "Prohibiciones y consecuencias para usuarios externos del Sistema Certificador INSO"
    doc.core_properties.subject = "Propuesta de capítulo reglamentario"
    doc.core_properties.author = "INSO"
    doc.core_properties.keywords = "INSO, sistema certificador, usuarios externos, prohibiciones, consecuencias, VUCE"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build_document()
